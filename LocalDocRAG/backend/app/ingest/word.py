import hashlib
import html
import io
import json
import mimetypes
import posixpath
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from uuid import UUID, uuid5

from lxml import etree
from PIL import Image

from app.core.errors import AppError

PARSER_VERSION = "ooxml-1.1.0"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "v": "urn:schemas-microsoft-com:vml"}
W = "{" + NS["w"] + "}"
R = "{" + NS["r"] + "}"


def hash_bytes(value):
    return hashlib.sha256(value).hexdigest()


def stable(version, path):
    return str(uuid5(UUID(version), path))


def safe_xml(value):
    if b"<!DOCTYPE" in value.upper() or b"<!ENTITY" in value.upper():
        raise AppError("UNSAFE_XML", "Word 包含不允许的 XML 实体。")
    return etree.fromstring(value, etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))


def validate_upload(data, suffix, settings):
    if len(data) > settings.max_upload:
        raise AppError("UPLOAD_TOO_LARGE", "文件超过大小限制。", 413)
    if suffix == ".doc":
        if not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise AppError("INVALID_DOC", "旧 Word 格式不正确或已加密。")
        return
    if suffix != ".docx" or not zipfile.is_zipfile(io.BytesIO(data)):
        raise AppError("INVALID_DOCX", "请选择真实 DOCX 或旧版 DOC 文件。")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > settings.max_entries or sum(e.file_size for e in entries) > settings.max_expanded:
            raise AppError("ZIP_LIMIT", "Word 展开大小或文件数量超限。")
        for entry in entries:
            path = entry.filename.replace("\\", "/")
            if path.startswith("/") or ".." in path.split("/") or ":" in path or entry.flag_bits & 1:
                raise AppError("UNSAFE_ARCHIVE", "Word 包含不安全路径或加密内容。")
            if entry.file_size > 1024 * 1024 and entry.file_size / max(entry.compress_size, 1) > 200:
                raise AppError("ZIP_RATIO_LIMIT", "Word 压缩比异常，拒绝展开。")
            if "vbaproject" in path.lower():
                raise AppError("MACRO_NOT_ALLOWED", "不支持包含宏的 Word 文件。")
        if "word/document.xml" not in archive.namelist():
            raise AppError("OOXML_MISSING", "缺少 Word 正文，可能是错误格式或加密文档。")


def libreoffice(source: Path, output: Path, fmt="docx"):
    import shutil
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise AppError("LIBREOFFICE_UNAVAILABLE", "需要安装 LibreOffice 才能转换旧 Word 或生成绘图回退预览。")
    with tempfile.TemporaryDirectory() as folder:
        profile = Path(folder) / "profile"
        profile.mkdir()
        (profile / "user").mkdir()
        (profile / "user/registrymodifications.xcu").write_text(
            '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>0</value></prop></item></oor:items>', encoding="utf-8")
        try:
            result = subprocess.run([executable, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--nologo", "--nodefault", "--norestore", "--convert-to", fmt, "--outdir", str(output), str(source)],
                                    timeout=90, capture_output=True)
        except subprocess.TimeoutExpired:
            raise AppError("DOC_CONVERSION_TIMEOUT", "Word 转换超时，原件已保留。") from None
        destination = output / (source.stem + "." + fmt)
        if result.returncode or not destination.is_file():
            raise AppError("DOC_CONVERSION_FAILED", "Word 转换失败；请检查是否受密码保护、已损坏或包含不支持的内容。")
        return destination


class WordParser:
    def __init__(self, settings):
        self.settings = settings

    def parse(self, file: Path, version: str, output: Path):
        output.mkdir(parents=True, exist_ok=True)
        original = file.read_bytes()
        validate_upload(original, file.suffix.lower(), self.settings)
        if file.suffix.lower() == ".doc":
            file = libreoffice(file, output)
            validate_upload(file.read_bytes(), ".docx", self.settings)
        # Also verifies the OOXML object model can open the package.
        from docx import Document
        try:
            Document(str(file))
        except Exception:
            raise AppError("CORRUPT_OOXML", "Word 对象结构损坏，原件已保留。") from None
        self.version, self.output = version, output
        self.blocks, self.tables, self.media, self.occurrences, self.warnings = [], [], {}, [], []
        self.headings = []
        with zipfile.ZipFile(file) as archive:
            self.archive = archive
            self.rels = {}
            rel_path = "word/_rels/document.xml.rels"
            if rel_path in archive.namelist():
                for rel in safe_xml(archive.read(rel_path)):
                    if rel.get("TargetMode") == "External":
                        self.warnings.append({"code": "EXTERNAL_RELATIONSHIP_BLOCKED", "relation": rel.get("Id")})
                        continue
                    target = posixpath.normpath(posixpath.join("word", rel.get("Target", "")))
                    if target.startswith("/"):
                        target = target.lstrip("/")
                    if ".." not in target.split("/"):
                        self.rels[rel.get("Id")] = target
            self.heading_styles = {}
            if "word/styles.xml" in archive.namelist():
                for style in safe_xml(archive.read("word/styles.xml")).findall("w:style", NS):
                    outline = style.find("w:pPr/w:outlineLvl", NS)
                    name = style.find("w:name", NS)
                    match = re.search(r"(?:heading|标题)\s*([1-9])", name.get(W + "val", "") if name is not None else "", re.I)
                    if outline is not None or match:
                        self.heading_styles[style.get(W + "styleId")] = int(outline.get(W + "val")) + 1 if outline is not None else int(match[1])
            body = safe_xml(archive.read("word/document.xml")).find("w:body", NS)
            if body is None:
                raise AppError("OOXML_BODY_MISSING", "Word 缺少正文。")
            self.walk(body, "body", None, None)
            media_parts = [n for n in archive.namelist() if n.startswith("word/media/") and not n.endswith("/")]
            for part in media_parts:
                if not any(part in m["parts"] for m in self.media.values()):
                    self.load_media(part, unreferenced=True)
            for part in [n for n in archive.namelist() if re.match(r"word/(charts|diagrams)/.+\.xml$", n)]:
                self.warnings.append({"code": "unsupported_visual", "part": part, "message": "图表或 SmartArt 需要渲染页级预览。"})
        if self.settings.docling:
            try:
                from docling.document_converter import DocumentConverter
                result = DocumentConverter().convert(str(file)).document
                (output / "docling.md").write_text(result.export_to_markdown(), encoding="utf-8")
                self.warnings.append({"code": "DOCLING_SUPPLEMENT", "message": "辅助结果保存在 docling.md；来源和顺序仍以 OOXML 为准。"})
            except Exception as exc:
                self.warnings.append({"code": "DOCLING_FAILED", "type": type(exc).__name__})
        if any(w["code"] == "unsupported_visual" for w in self.warnings):
            self.render_fallback(file)
        canonical = {"parser_version": PARSER_VERSION, "blocks": self.blocks, "tables": self.tables,
                     "media": list(self.media.values()), "media_occurrences": self.occurrences}
        supported_parts = {p for m in self.media.values() if m["status"] != "unsupported" for p in m["parts"]}
        report = {"parser_version": PARSER_VERSION, "media_parts": len(media_parts), "processed_media_parts": len(supported_parts),
                  "unsupported_media_parts": len(media_parts) - len(supported_parts), "extracted_images": len(self.media),
                  "image_occurrences": len(self.occurrences), "tables": len(self.tables), "blocks": len(self.blocks),
                  "sections": sum(b["type"] == "heading" for b in self.blocks), "warnings": self.warnings, "ocr_success": 0}
        (output / "canonical_document.json").write_text(json.dumps(canonical, ensure_ascii=False, indent=2), encoding="utf-8")
        (output / "parse_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return canonical, report

    def block(self, kind, text, anchor, parent=None, structured=None):
        item = {"id": stable(self.version, anchor), "version_id": self.version, "parent_block_id": parent,
                "block_order": len(self.blocks), "type": kind, "heading_path": list(self.headings), "anchor": anchor,
                "plain_text": text, "structured": structured or {}, "warnings": []}
        self.blocks.append(item)
        return item

    def walk(self, element, path, parent, position):
        counts = {}
        for child in element:
            tag = etree.QName(child).localname
            counts[tag] = counts.get(tag, 0) + 1
            anchor = f"{path}/{tag}[{counts[tag]}]"
            if child.tag == W + "p":
                self.paragraph(child, anchor, parent, position)
            elif child.tag == W + "tbl":
                self.table(child, anchor, parent)
            elif tag in {"sdt", "sdtContent", "customXml"}:
                self.walk(child, anchor, parent, position)

    def paragraph(self, paragraph, path, parent, position):
        text = "".join(paragraph.xpath(".//w:t/text() | .//w:tab/text()", namespaces=NS))
        style = paragraph.find("w:pPr/w:pStyle", NS)
        outline = paragraph.find("w:pPr/w:outlineLvl", NS)
        level = int(outline.get(W + "val")) + 1 if outline is not None else self.heading_styles.get(style.get(W + "val") if style is not None else "")
        if level and parent is None:
            self.headings = self.headings[:level - 1] + [text]
        listing = paragraph.find("w:pPr/w:numPr/w:ilvl", NS)
        block = self.block("heading" if level else "paragraph", text, path, parent,
                           {"heading_level": level, "list_level": int(listing.get(W + "val", "0")) if listing is not None else None, "inline": []})
        style_id = style.get(W + "val", "") if style is not None else ""
        if re.search(r"^(toc\d*|目录\d*)$", style_id, re.I) or paragraph.xpath(".//w:instrText[contains(., 'PAGEREF')]", namespaces=NS):
            block["structured"]["navigation"] = True
        for node in paragraph.iter():
            if node.tag == W + "t":
                block["structured"]["inline"].append({"type": "text", "text": node.text or ""})
            elif etree.QName(node).localname in {"blip", "imagedata"}:
                rid = node.get(R + "embed") or node.get(R + "id")
                part = self.rels.get(rid)
                if not part or part not in self.archive.namelist():
                    self.warnings.append({"code": "IMAGE_RELATIONSHIP_MISSING", "anchor": path})
                    continue
                media = self.load_media(part)
                occurrence = {"id": stable(self.version, path + f"/image[{len(block['structured']['inline'])}]"), "media_id": media["id"],
                              "block_id": block["id"], "table_id": position[0] if position else None,
                              "row_index": position[1] if position else None, "col_index": position[2] if position else None,
                              "occurrence_order": len(self.occurrences), "anchor": path}
                self.occurrences.append(occurrence)
                block["structured"]["inline"].append({"type": "image", "media_id": media["id"], "occurrence_id": occurrence["id"]})
            elif etree.QName(node).localname in {"chart", "relIds", "OLEObject"}:
                self.warnings.append({"code": "unsupported_visual", "anchor": path, "message": "发现无法直接解析的绘图或嵌入对象。"})

    def load_media(self, part, unreferenced=False):
        data = self.archive.read(part)
        digest = hash_bytes(data)
        media_id = stable(self.version, "media:" + digest)
        if media_id in self.media:
            if part not in self.media[media_id]["parts"]:
                self.media[media_id]["parts"].append(part)
            return self.media[media_id]
        suffix = Path(part).suffix.lower()
        destination = self.settings.data / "media" / (digest + suffix)
        destination.write_bytes(data)
        mime = mimetypes.guess_type(part)[0] or "application/octet-stream"
        media = {"id": media_id, "version_id": self.version, "sha256": digest, "storage_path": str(destination),
                 "mime": mime, "width": None, "height": None, "status": "pending", "parts": [part], "warnings": [], "ocr": None, "vision": None}
        self.media[media_id] = media
        if len(self.media) > self.settings.max_images:
            media["status"] = "unsupported"
            media["warnings"].append("IMAGE_COUNT_LIMIT")
        else:
            try:
                with Image.open(io.BytesIO(data)) as image:
                    media["width"], media["height"] = image.size
                    if image.width * image.height > self.settings.max_pixels:
                        raise ValueError("pixel limit")
                    if suffix in {".emf", ".wmf"}:
                        # Pillow uses the local Windows metafile renderer. Keep
                        # the original Word/part and expose a bounded PNG for
                        # OCR and the browser; other platforms retain fallback.
                        image.thumbnail((self.settings.max_image_side, self.settings.max_image_side))
                        raster = image.convert("RGB")
                        rendered = self.settings.data / "media" / (digest + ".raster.png")
                        raster.save(rendered, "PNG")
                        media.update({"storage_path": str(rendered), "mime": "image/png", "width": raster.width, "height": raster.height})
                        media["warnings"].append("VECTOR_IMAGE_RASTERIZED")
                    else:
                        image.verify()
                if media["mime"] not in {"image/png", "image/jpeg", "image/gif", "image/webp", "image/bmp", "image/tiff"}:
                    raise ValueError("unsupported image")
            except Exception:
                media["status"] = "unsupported"
                media["warnings"].append("unsupported_visual")
        if unreferenced:
            media["warnings"].append("UNREFERENCED_MEDIA_PART")
        for warning in media["warnings"]:
            self.warnings.append({"code": warning, "media_id": media_id, "part": part})
        return media

    def table(self, table, path, parent):
        block = self.block("table", "", path, parent)
        table_id = stable(self.version, path + ":table")
        cells, active = [], {}
        rows = table.findall("w:tr", NS)
        for row_index, row in enumerate(rows):
            before = row.find("w:trPr/w:gridBefore", NS)
            col = int(before.get(W + "val", "0")) if before is not None else 0
            next_active = {}
            for cell_index, cell in enumerate(row.findall("w:tc", NS)):
                grid = cell.find("w:tcPr/w:gridSpan", NS)
                span = int(grid.get(W + "val", "1")) if grid is not None else 1
                merge = cell.find("w:tcPr/w:vMerge", NS)
                if merge is not None and merge.get(W + "val", "continue") == "continue" and col in active:
                    origin = active[col]
                    if origin["colspan"] != span:
                        self.warnings.append({"code": "MERGE_SPAN_CONFLICT", "anchor": path, "row": row_index, "col": col})
                    origin["rowspan"] += 1
                    for c in range(col, col + span):
                        next_active[c] = origin
                    # Preserve unexpected continuation text/images rather than silently drop it.
                    if cell.xpath(".//w:t | .//a:blip", namespaces=NS):
                        self.warnings.append({"code": "MERGE_CONTINUATION_CONTENT", "anchor": path, "row": row_index, "col": col})
                        self.walk(cell, f"{path}/tr[{row_index+1}]/tc[{cell_index+1}]", block["id"], (table_id, row_index, col))
                    col += span
                    continue
                start = len(self.blocks)
                anchor = f"{path}/tr[{row_index+1}]/tc[{cell_index+1}]"
                self.walk(cell, anchor, block["id"], (table_id, row_index, col))
                refs = [b["id"] for b in self.blocks[start:] if b["parent_block_id"] == block["id"]]
                value = "\n".join(b["plain_text"] for b in self.blocks[start:] if b["type"] != "table")
                item = {"row": row_index, "col": col, "rowspan": 1, "colspan": span, "text": value,
                        "block_ids": refs, "anchor": anchor, "header": row.find("w:trPr/w:tblHeader", NS) is not None or row_index == 0}
                cells.append(item)
                if merge is not None:
                    for c in range(col, col + span):
                        next_active[c] = item
                col += span
            active = next_active
        structure = {"id": table_id, "rows": len(rows), "cols": max((c["col"] + c["colspan"] for c in cells), default=0), "cells": cells}
        rendered = "<table>" + "".join("<tr>" + "".join(f'<{"th" if c["header"] else "td"} rowspan="{c["rowspan"]}" colspan="{c["colspan"]}">{html.escape(c["text"])}</{"th" if c["header"] else "td"}>' for c in cells if c["row"] == r) + "</tr>" for r in range(len(rows))) + "</table>"
        block["structured"] = {"table_id": table_id}
        block["plain_text"] = "\n".join(c["text"] for c in cells)
        self.tables.append({"id": table_id, "version_id": self.version, "block_id": block["id"], "caption": "表格 " + str(len(self.tables) + 1), "table_json": structure, "sanitized_html": rendered})

    def render_fallback(self, file):
        try:
            pdf = libreoffice(file, self.output, "pdf")
            import shutil
            renderer = shutil.which("pdftoppm")
            if not renderer:
                raise AppError("PDF_RENDERER_UNAVAILABLE", "缺少 Poppler，PDF 已保留。")
            subprocess.run([renderer, "-scale-to", "1400", "-png", str(pdf), str(self.output / "rendered")], timeout=90, check=True, capture_output=True)
            for index, page in enumerate(sorted(self.output.glob("rendered-*.png")), 1):
                data = page.read_bytes()
                digest = hash_bytes(data)
                destination = self.settings.data / "media" / (digest + ".png")
                destination.write_bytes(data)
                mid = stable(self.version, "rendered:" + str(index))
                with Image.open(page) as im:
                    width, height = im.size
                self.media[mid] = {"id": mid, "version_id": self.version, "sha256": digest, "storage_path": str(destination), "mime": "image/png", "width": width, "height": height,
                                   "status": "pending", "parts": [], "warnings": ["LIBREOFFICE_RENDERED_PAGE"], "ocr": None, "vision": None}
                b = self.block("rendered_page", "LibreOffice 排版页 " + str(index), "rendered-page:" + str(index), structured={"media_id": mid, "rendered_page": index, "precision": "rendered_page"})
                self.occurrences.append({"id": stable(self.version, b["anchor"]), "media_id": mid, "block_id": b["id"], "table_id": None, "row_index": None, "col_index": None, "occurrence_order": len(self.occurrences), "anchor": b["anchor"]})
        except Exception as exc:
            self.warnings.append({"code": "RENDER_FALLBACK_FAILED", "type": type(exc).__name__, "message": "不支持的绘图已保留记录，请下载原件查看。"})
