import asyncio
import io

from docx import Document
from fastapi.testclient import TestClient

from app.db.database import uid
from app.ingest.word import WordParser
from app.main import create_app
from app.rag.chunks import make_chunks
from app.rag.models import Ollama
from app.rag.service import RAG, PROCEDURE_SECTIONS, answer_sources
from app.db.database import Database, dump, now
from app.core.errors import AppError
import pytest
from app.core.model_queue import ModelQueue, process_alive
import os
import time
from test_localdoc import FakeModels, FakeVectors, fixtures, settings


def test_answer_evidence_keeps_cross_table_completion_conditions():
    approval = {'type':'table_row','table_id':'request','fs_ids':['FS-1-009']}
    request_table = {'type':'operation','table_id':'request','fs_ids':['FS-1-009']}
    plan_table = {'type':'operation','table_id':'plan','fs_ids':['FS-2-010']}
    completion = {'type':'table_row','table_id':'plan','fs_ids':['FS-2-010']}
    sources = [approval, request_table, plan_table, completion]
    assert answer_sources(sources, '审批后代表任务完成了吗？') == [approval, completion]
    assert answer_sources(sources, 'FS-2-010 什么时候完成？') == [completion]
    # A full-process question still sees the complete original table.
    assert answer_sources([request_table, approval, completion], '完整流程步骤') == [request_table, completion]


def test_procedure_field_instructions_and_citations_reach_the_model(settings):
    db=Database(settings.database);db.migrate();kid=uid()
    db.execute('INSERT INTO knowledge_bases VALUES(?,?,?,?,?,NULL)', (kid,'test','',now(),now()))
    class ProcedureModels(FakeModels):
        async def chat(self, messages, schema=None, **kwargs):
            assert 'reviewers' in schema['required']
            assert '审核/审批区块' in messages[0]['content']
            assert schema['$defs']['AnswerParagraph']['properties']['source_ids']['items']['enum']==['S1']
            return dump({**{key:{'text':'原文约束','source_ids':['S1']} for key in PROCEDURE_SECTIONS},'insufficient':False})
    rag=RAG(settings,db,ProcedureModels(),FakeVectors())
    async def search(*args):
        return {'fused':[{'source_id':'S1','type':'text','section_path':[],'content':'审核/审批区块：审核人必填、可多选','warnings':[]}],'warnings':[]}
    rag.search=search
    result=asyncio.run(rag.answer(kid,'如何办理完整流程？'))
    assert '审核与审批人员：原文约束[S1]' in result['answer']
    assert len(result['sources'])==1 and not result['insufficient']


def test_partial_model_response_is_never_a_final_answer(settings):
    class PartialModel(Ollama):
        async def serialized(self, model, path, payload, priority):
            assert payload['options']['num_predict'] == 2048
            return {'done':False,'message':{'content':'','thinking':'partial'}}
    async def run():
        model=PartialModel(settings,None)
        try:
            with pytest.raises(AppError,match='输出预算'):
                await model.chat([{'role':'user','content':'test'}])
        finally:
            await model.close()
    asyncio.run(run())


def test_dead_queue_owner_does_not_block_a_restart(settings):
    from app.db.database import Database
    db=Database(settings.database); db.migrate()
    dead_owner='2147483647:dead'
    assert process_alive(f'{os.getpid()}:live')
    assert not process_alive(dead_owner)
    db.execute('INSERT INTO model_queue VALUES(?,?,?,?,?)',('old',dead_owner,0,time.time()-1,time.time()+720))
    async def run():
        async with ModelQueue(db,2).acquire(10):
            assert not db.one("SELECT id FROM model_queue WHERE id='old'")
            assert db.one("SELECT owner FROM leases WHERE name='gpu'")
    asyncio.run(run())


def test_complete_row_keeps_tail_and_repeats_identity(settings, fixtures):
    version = uid()
    canonical, _ = WordParser(settings).parse(fixtures / "training_manual.docx", version, settings.data / "parsed")
    table = canonical["tables"][0]
    cells = table["table_json"]["cells"]
    target = next(c for c in cells if c["row"] == 2 and c["col"] == 2)
    target["text"] = "条件说明" * 180 + "末尾必须检查审批节点"
    chunks = make_chunks(canonical, version, uid(), "test.docx")
    row = next(c for c in chunks if c["type"] == "table_row" and c["source_meta"]["row"] == 2)
    assert "末尾必须检查审批节点" in row["content"]
    assert "管理人员" in row["content"]


def test_operation_overview_keeps_source_fs_ids_and_skips_navigation(settings, fixtures):
    version = uid()
    canonical, _ = WordParser(settings).parse(fixtures / "training_manual.docx", version, settings.data / "parsed")
    table = canonical["tables"][0]
    cells = table["table_json"]["cells"]
    next(c for c in cells if c["row"] == 1 and c["col"] == 0)["text"] = "FS-5.4.2.7-001"
    next(c for c in cells if c["row"] == 2 and c["col"] == 0)["text"] = "FS-5.4.2.7-008"
    chunks = make_chunks(canonical, version, uid(), "test.docx")
    overview = next(c for c in chunks if c["type"] == "operation")
    assert overview["source_meta"]["fs_ids"] == ["FS-5.4.2.7-001", "FS-5.4.2.7-008"]
    assert overview["source_meta"]["row"] is None
    canonical["blocks"][0]["structured"]["navigation"] = True
    chunks2 = make_chunks(canonical, version, uid(), "test.docx")
    assert all(c["block_refs"] != [canonical["blocks"][0]["id"]] for c in chunks2)


def test_agent_reads_complete_operation_chain_and_exact_fs_beats_generic_rows(settings, fixtures):
    word = Document(fixtures / "training_manual.docx")
    word.tables[0].cell(1, 0).text = "FS-5.4.2.7-001"
    word.tables[0].cell(1, 2).text = "前置条件" * 40 + "岗位受训人必须选择"
    word.tables[0].cell(2, 0).text = "FS-5.4.2.7-008"
    word.tables[0].cell(2, 2).text = "点击发起需要电子签名，提交成功后进入部门审批"
    buffer = io.BytesIO()
    word.save(buffer)
    app = create_app(settings, FakeModels(), FakeVectors())
    with TestClient(app, headers={"x-localdoc-request":"1"}) as client:
        kid = client.post('/api/knowledge-bases',json={"name":"实际操作结构测试"}).json()['id']
        item = client.post(f'/api/knowledge-bases/{kid}/documents',files={'files':('manual.docx',buffer.getvalue())}).json()['items'][0]
        asyncio.run(app.state.worker.run_job(app.state.db.one('SELECT * FROM jobs WHERE id=?',(item['job_id'],))))
        hits = client.post('/api/agent/search',json={'knowledge_base_id':kid,'query':'FS-5.4.2.7-008'}).json()['results']
        assert 'FS-5.4.2.7-008' in hits[0]['snippet']
        cid = hits[0]['id'].removeprefix('localdoc:')
        read = client.get(f'/api/agent/knowledge/{cid}?knowledge_base_id={kid}').json()
        assert '部门审批' in read['content']
        overview = app.state.db.one("SELECT id FROM chunks WHERE version_id=? AND type='operation'",(item['version_id'],))
        assert 'localdoc:'+overview['id'] in read['content']
        full = client.get(f"/api/agent/knowledge/{overview['id']}?knowledge_base_id={kid}").json()
        assert '岗位受训人必须选择' in full['content']
        assert 'FS-5.4.2.7-001' in full['content'] and 'FS-5.4.2.7-008' in full['content']
        assert full['metadata']['source']['operation_complete'] is True
