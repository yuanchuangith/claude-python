import asyncio
import hashlib
import io
import json

from docx import Document
from fastapi.testclient import TestClient

from app.main import create_app
from test_localdoc import FakeModels, FakeVectors, fixtures, settings


class EvidenceModels(FakeModels):
    async def tags(self):
        return [{"name": "qwen3-embedding:0.6b", "digest": "test-digest"}]

    async def chat(self, *args, **kwargs):
        raise AssertionError("Agent evidence access must not invoke local chat")

    async def vision(self, *args, **kwargs):
        raise AssertionError("Reading indexed evidence must not invoke vision")


def ingest(client, app, body, name="manual.docx"):
    kid = client.post("/api/knowledge-bases", json={"name": "共享手册"}).json()["id"]
    item = client.post(f"/api/knowledge-bases/{kid}/documents", files={"files": (name, body)}).json()["items"][0]
    asyncio.run(app.state.worker.run_job(app.state.db.one("SELECT * FROM jobs WHERE id=?", (item["job_id"],))))
    assert app.state.db.one("SELECT status FROM document_versions WHERE id=?", (item["version_id"],))["status"].startswith("completed")
    return kid, item


def test_shared_evidence_provenance_complete_non_fs_table_and_version_fence(settings, fixtures):
    body = (fixtures / "training_manual.docx").read_bytes()
    app = create_app(settings, EvidenceModels(), FakeVectors())
    with TestClient(app, headers={"x-localdoc-request": "1"}) as client:
        kid, item = ingest(client, app, body)
        health = client.get("/api/health").json()
        assert health["ready"] and health["evidence_ready"]
        assert not health["chat_ready"] and not health["image_caption_ready"]
        query = {"knowledge_base_id": kid, "query": "管理人员", "limit": 1}
        search = client.post("/api/agent/search", json=query).json()
        assert search["results"]
        continuation = client.post("/api/agent/search", json={**query, "offset": 1, "result_version": search["resultVersion"]})
        assert continuation.status_code == 200
        cid = search["results"][0]["id"].removeprefix("localdoc:")
        read = client.get(f"/api/agent/knowledge/{cid}?knowledge_base_id={kid}").json()
        meta = read["metadata"]
        assert meta["contentPurpose"] == "business_knowledge" and meta["status"] == "evidence_only"
        assert meta["consumerAgentIds"] == ["action_design", "business_operation"]
        source = meta["source"]
        assert source["document_sha256"] == hashlib.sha256(body).hexdigest()
        assert source["version_id"] == item["version_id"] and source["document_id"] == item["document_id"]
        assert source["sha256"] != source["document_sha256"]
        assert source["section_path"] and source["anchor"] and source["preview_url"]
        assert meta["fullTableId"] in read["content"]
        assert search["results"][0]["fullTableId"] == meta["fullTableId"]
        assert "knowledge_read" in search["results"][0]["readHint"] and "nextOffset" in search["results"][0]["readHint"]
        tid = meta["fullTableId"].removeprefix("localdoc:")
        full = client.get(f"/api/agent/knowledge/{tid}?knowledge_base_id={kid}").json()
        assert full["metadata"]["source"]["table_complete"] is True
        assert "管理人员" in full["content"] and "90分" in full["content"] and "合格分数" in full["content"]
        other = client.post("/api/knowledge-bases", json={"name": "另一知识库"}).json()["id"]
        assert client.get(f"/api/agent/knowledge/{tid}?knowledge_base_id={other}").json() is None
        app.state.db.execute("UPDATE documents SET active_version_id=NULL WHERE id=?", (item["document_id"],))
        changed = client.post("/api/agent/search", json={**query, "offset": 1, "result_version": search["resultVersion"]})
        assert changed.status_code == 409 and changed.json()["code"] == "KNOWLEDGE_SEARCH_CHANGED"
        assert client.get(f"/api/agent/knowledge/{cid}?knowledge_base_id={kid}").json() is None
        assert client.get(f"/api/agent/knowledge/{tid}?knowledge_base_id={kid}").json() is None


def test_search_table_hints_keep_exact_row_and_choose_existing_operation_or_canonical_table(settings):
    word = Document()
    word.add_heading("来源流程", level=1)
    operation = word.add_table(rows=2, cols=3)
    for col, value in enumerate(["FS-ID", "操作", "描述"]):
        operation.cell(0, col).text = value
    for col, value in enumerate(["FS-4.2-001", "发起", "选择资料编号后发起"]):
        operation.cell(1, col).text = value
    word.add_heading("字段信息", level=2)
    fields = word.add_table(rows=3, cols=2)
    for row, values in enumerate([("字段", "必填项"), ("资料编号", "是"), ("说明", "否")]):
        for col, value in enumerate(values):
            fields.cell(row, col).text = value
    buffer = io.BytesIO()
    word.save(buffer)
    app = create_app(settings, EvidenceModels(), FakeVectors())
    with TestClient(app, headers={"x-localdoc-request": "1"}) as client:
        kid, item = ingest(client, app, buffer.getvalue())
        result = client.post("/api/agent/search", json={"knowledge_base_id": kid, "query": "资料编号", "limit": 20}).json()
        rows = [hit for hit in result["results"] if "table_row" in hit["tags"] and "资料编号" in hit["snippet"]]
        assert len(rows) == 2
        target_types = set()
        for hit in rows:
            row = client.get(f"/api/agent/knowledge/{hit['id'].removeprefix('localdoc:')}?knowledge_base_id={kid}").json()
            assert row["id"] == hit["id"] and row["metadata"]["sourceVersion"] == item["version_id"]
            assert hit["fullTableId"] == row["metadata"]["fullTableId"] and hit["fullTableId"] != hit["id"]
            target = client.get(f"/api/agent/knowledge/{hit['fullTableId'].removeprefix('localdoc:')}?knowledge_base_id={kid}").json()
            assert target["metadata"]["sourceVersion"] == item["version_id"]
            assert target["metadata"]["source"]["table_id"] == row["metadata"]["source"]["table_id"]
            assert target["metadata"]["source"]["table_complete"] is True
            assert target["metadata"]["status"] == "evidence_only"
            target_types.add(target["metadata"]["source"]["operation_complete"])
        assert target_types == {True, False}


def test_search_rejects_activation_during_embedding(settings, fixtures):
    models = EvidenceModels()
    app = create_app(settings, models, FakeVectors())
    with TestClient(app, headers={"x-localdoc-request": "1"}) as client:
        kid, item = ingest(client, app, (fixtures / "training_manual.docx").read_bytes())
        original_embed = models.embed

        async def switch_version(*args, **kwargs):
            app.state.db.execute("UPDATE documents SET active_version_id=NULL WHERE id=?", (item["document_id"],))
            return await original_embed(*args, **kwargs)

        models.embed = switch_version
        response = client.post("/api/agent/search", json={"knowledge_base_id": kid, "query": "管理人员"})
        assert response.status_code == 409 and response.json()["code"] == "KNOWLEDGE_VERSION_CHANGED"


def test_adjacent_explanations_keep_sibling_prose_without_repeating_tables_or_descendants(settings):
    word = Document()
    word.add_heading("通用操作", level=1)
    word.add_paragraph("办理前说明")
    operation = word.add_table(rows=2, cols=3)
    for row, values in enumerate([("FS-ID", "操作", "描述"), ("FS-9.1-001", "唯一操作正文", "选择资料")]):
        for col, value in enumerate(values):
            operation.cell(row, col).text = value
    nested = operation.cell(1, 2).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套表后代"
    word.add_paragraph("办理后说明")
    fields = word.add_table(rows=2, cols=2)
    for row, values in enumerate([("字段", "必填项"), ("唯一字段正文", "是")]):
        for col, value in enumerate(values):
            fields.cell(row, col).text = value
    barrier = word.add_table(rows=1, cols=1)
    barrier.cell(0, 0).text = "独立相邻表不能当说明"
    word.add_paragraph("另一表之后的远处说明")
    buffer = io.BytesIO()
    word.save(buffer)
    body = buffer.getvalue()
    app = create_app(settings, EvidenceModels(), FakeVectors())
    with TestClient(app, headers={"x-localdoc-request": "1"}) as client:
        kid, item = ingest(client, app, body)
        chunk = app.state.db.one("SELECT * FROM chunks WHERE version_id=? AND type='operation'", (item["version_id"],))
        table_id = json.loads(chunk["source_meta_json"])["table_id"]
        for cid in [chunk["id"], table_id]:
            read = client.get(f"/api/agent/knowledge/{cid}?knowledge_base_id={kid}").json()
            source = read["metadata"]["source"]
            assert source["table_complete"] is True and source["version_id"] == item["version_id"]
            assert source["document_sha256"] == hashlib.sha256(body).hexdigest()
            assert read["content"].count("FS-9.1-001") == 1
            assert read["content"].count("唯一操作正文") == 1
            assert read["content"].count("嵌套表后代") == 1
            adjacent = json.loads(read["content"].split("\n相邻说明：", 1)[1])
            assert [block["plain_text"] for block in adjacent] == ["办理前说明", "办理后说明"]
        row = app.state.db.one("SELECT * FROM chunks WHERE version_id=? AND type='table_row' AND instr(content,'FS-9.1-001')>0", (item["version_id"],))
        assert [b["plain_text"] for b in app.state.rag.hydrate(row)["adjacent_blocks"]] == ["办理前说明", "办理后说明"]
        field_table = app.state.db.one("SELECT * FROM tables WHERE version_id=? AND instr(table_json,'唯一字段正文')>0", (item["version_id"],))
        read = client.get(f"/api/agent/knowledge/{field_table['id']}?knowledge_base_id={kid}").json()
        assert read["content"].count("唯一字段正文") == 1
        adjacent = json.loads(read["content"].split("\n相邻说明：", 1)[1])
        assert [block["plain_text"] for block in adjacent] == ["办理后说明"]
        assert "独立相邻表不能当说明" not in read["content"] and "远处说明" not in read["content"]
        assert app.state.db.one("SELECT source_meta_json FROM chunks WHERE id=?", (chunk["id"],))["source_meta_json"] == chunk["source_meta_json"]


def test_cross_chapter_conflicting_source_branches_remain_separate(settings):
    word = Document()
    for heading, rule in [("申请条件", "满足条件甲时选择矩阵资料"), ("例外条件", "满足条件甲时选择风险内容")]:
        word.add_heading(heading, level=1)
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "编号", "规则"
        table.cell(1, 0).text, table.cell(1, 1).text = "FS-1.2-003", rule
    buffer = io.BytesIO()
    word.save(buffer)
    app = create_app(settings, EvidenceModels(), FakeVectors())
    with TestClient(app, headers={"x-localdoc-request": "1"}) as client:
        kid, _ = ingest(client, app, buffer.getvalue())
        search = client.post("/api/agent/search", json={"knowledge_base_id": kid, "query": "FS-1.2-003", "limit": 20}).json()
        reads = [client.get(f"/api/agent/knowledge/{hit['id'].removeprefix('localdoc:')}?knowledge_base_id={kid}").json() for hit in search["results"]]
        clauses = [read for read in reads if read["metadata"]["source"]["row"] == 1]
        assert len(clauses) == 2
        assert {tuple(read["metadata"]["source"]["section_path"]) for read in clauses} == {("申请条件",), ("例外条件",)}
        assert any("选择矩阵资料" in read["content"] for read in clauses)
        assert any("选择风险内容" in read["content"] for read in clauses)
        assert all(read["metadata"]["status"] == "evidence_only" for read in clauses)
